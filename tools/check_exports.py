"""check_exports.py — recalculate every formula in the Excel exports (with pycel) and compare with the model.
Run after tools/sample_data.py:  pip install pycel; python3 tools/check_exports.py"""
import sys, json, re
sys.path.insert(0, "tools")
import winutf8; winutf8.ensure()   # Windows: run in UTF-8 mode (the tools write characters like ¢ and ▲)
from pycel import ExcelCompiler
import openpyxl
import financial_model as fm, deliveries as dl
import exportkit as ek

res = fm.build()
bad = 0
for org in ["trades", "services", "nfp"]:
    path = f"media/exports/{org}-sample-statements.xlsx"
    xc = ExcelCompiler(filename=path)
    wb = openpyxl.load_workbook(path)
    n = 0
    # statements: every row vs the model
    out = res[org]["out"]
    for tab, key in (("Income and expenditure" if org == "nfp" else "Profit and loss", "pnl"), ("Balance sheet", "bsr"), ("Cash flow", "cfr")):
        ws = wb[tab]
        rows = [rw for rw in out[key]]
        r = 5
        for rw in rows:
            if rw["values"] is not None:
                for c, col in ((0, "B"), (1, "C")):
                    v = xc.evaluate(f"'{tab}'!{col}{r}")
                    n += 1
                    if round(v) != rw["values"][c]:
                        bad += 1; print("MISMATCH", org, tab, rw["label"], v, rw["values"][c])
                ch = xc.evaluate(f"'{tab}'!D{r}")
                if round(ch) != rw["values"][0] - rw["values"][1]:
                    bad += 1; print("CHANGE", org, tab, rw["label"])
            r += 1
    # fourth sheet KPIs vs model values
    f4 = res[org]["fourth"]
    for i, k in enumerate(f4["kpis"]):
        exp = k["support"]["xl"]["rows"][-1]["values"]
        got = xc.evaluate(f"'The fourth sheet'!B{6 + i}")
        n += 1
        if abs(got - exp[0]) > 1e-9:
            bad += 1; print("KPI", org, k["label"], got, exp[0])
    # every formula cell evaluates without error
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    v = xc.evaluate(f"'{ws.title}'!{cell.coordinate}")
                    n += 1
                    if v is None or (isinstance(v, str) and v.startswith("#")):
                        if not (isinstance(v, str) and v == ""):
                            bad += 1; print("ERR", org, ws.title, cell.coordinate, cell.value[:60], v)
    print(org, "checked", n)
# deliveries summary vs payload KPIs
d = dl.build()
xc = ExcelCompiler(filename="media/exports/deliveries-sample.xlsx")
r = 7
for side in ("out", "in"):
    for p in d["payload"]["periods"]:
        k = d["payload"][side][p["id"]]["kpis"]
        got_total = xc.evaluate(f"Summary!D{r}")
        got_ot = xc.evaluate(f"Summary!J{r}")
        if got_total != k["total"] or (k["on_time_pct"] is not None and abs(got_ot * 100 - k["on_time_pct"]) > 0.051):
            bad += 1; print("DELIV", side, p["id"], got_total, k["total"], got_ot, k["on_time_pct"])
        r += 1
    r += 4
print("deliveries summary checked")

# every report, every period: formulas calculate, and each headline number in Excel = the number on the site
from fourthsheet.build import load_data, periods_for, run_report, out_dir
from fourthsheet.catalogue import REPORTS
import build_reports as br
D = load_data()
nrep = nf = 0
for slug in REPORTS:
    for per in periods_for(slug, D):
        for r in br.file_parts(run_report(slug, per, D)):          # one workbook per business
            path = str(out_dir(r) / f"{br.file_name(r)}.xlsx")
            xc = ExcelCompiler(filename=path)
            wb = openpyxl.load_workbook(path)
            want = {}
            for b in r["blocks"]:
                for sec in b["sections"]:
                    secs = list(sec["by"].values()) if sec["type"] == "vary" else [sec]
                    for x in secs:
                        if x["type"] == "kpis":
                            for k in x["items"]:      # the figure as the site shows it (never "the last line of its workings")
                                if ek.shown(k["value"]) is not None:
                                    want.setdefault(k["label"], []).append(ek.shown(k["value"]))
            ws = wb["Report"]
            for row in ws.iter_rows(min_col=1, max_col=2):
                lab, cell = row[0].value, row[1]
                if lab in want and isinstance(cell.value, str) and cell.value.startswith("=Workings!"):
                    got = xc.evaluate(f"Report!{cell.coordinate}")
                    exp, tol = want[lab].pop(0) if want[lab] else (None, 0)
                    nf += 1
                    if exp is not None and (not isinstance(got, (int, float)) or abs(got - exp) > tol):
                        bad += 1; print("REPORT KPI", slug, per, lab, got, exp)
            for w in wb.worksheets:
                for row in w.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and cell.value.startswith("="):
                            v = xc.evaluate(f"'{w.title}'!{cell.coordinate}")
                            nf += 1
                            if v is None or (isinstance(v, str) and v.startswith("#")):
                                bad += 1; print("REPORT ERR", slug, per, w.title, cell.coordinate, cell.value[:60], v)
            nrep += 1
print(f"reports: {nrep} workbooks, {nf} formulas and headline numbers checked")
print("PROBLEMS:", bad)
