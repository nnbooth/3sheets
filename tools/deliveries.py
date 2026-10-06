#!/usr/bin/env python3
"""
deliveries.py — mock deliveries in and out for a Brisbane distribution centre,
for the map example on examples.html.

Sample Distribution's DC is at Wacol. Every delivery OUT to a customer
(~65 customers across south-east Queensland) and every delivery IN from a
supplier (around Australia, plus the Port of Brisbane and air freight) is one
row, with promised and actual dates, carrier, cartons, value and the reason
when it was late. Daily, from 3 August 2026 to today (Tuesday 6 October 2026,
2pm), so the map can show today, the last working day, the last 7 or 30 days.

Stories in the data (what the map should make obvious):
  - From 7 September, Carrier A's Gold Coast run starts running late.
  - Supplier C (Melbourne) runs late through September: line-haul delays.
  - One import container was held at the Port of Brisbane in late September.

Writes (via tools/sample_data.py):
  deliveries-data.js                       what the map needs (aggregated)
  media/exports/deliveries-sample.xlsx     formatted workbook (every delivery)
  media/exports/deliveries-sample.pdf      A4 landscape report with maps
  CSV tables: dim_customer, dim_supplier, fact_delivery_out, fact_delivery_in
All data is invented.
"""

import json
import random
from datetime import date, datetime, timedelta
from pathlib import Path

import data_status as ds
import exportkit as ek
from warehouse import QLD_HOLIDAYS, key, working


def period_status(d_from, d_to):
    """How final a date range is: the least final day decides the badge."""
    days_, d = [], d_from
    while d <= d_to:
        days_.append(ds.day_status(d))
        d += timedelta(days=1)
    notes = []
    if "Incomplete" in days_:
        notes.append(f"today is incomplete (as at {ds.as_at_text().split(',')[0]}; deliveries still on the road)")
    if "Provisional" in days_:
        notes.append(f"October's earlier days are provisional until it locks on {ds.lock_date('2026-10').strftime('%-d %b')} (late proofs of delivery can still change them)")
    if "Locked" in days_ and notes:
        notes.append("earlier days are locked")
    badge = "Incomplete" if "Incomplete" in days_ else "Provisional" if "Provisional" in days_ else "Locked"
    text = "; ".join(notes) if notes else "every day in this period is locked"
    return badge, text[0].upper() + text[1:] + "."

REPO = Path(__file__).resolve().parent.parent
EXPORTS = REPO / "media" / "exports"
START, TODAY = date(2026, 8, 3), date(2026, 10, 6)
NOW_HOUR = 14   # "as at 2pm today"
DC = {"name": "Sample Distribution DC, Wacol", "lat": -27.588, "lon": 152.933}

# suburb, latitude, longitude, region
SUBURBS = [
    ("Brisbane CBD", -27.4698, 153.0251, "Inner Brisbane"), ("Fortitude Valley", -27.4570, 153.0346, "Inner Brisbane"),
    ("Newstead", -27.4467, 153.0450, "Inner Brisbane"), ("South Brisbane", -27.4800, 153.0180, "Inner Brisbane"),
    ("West End", -27.4820, 153.0090, "Inner Brisbane"), ("Woolloongabba", -27.4890, 153.0360, "Inner Brisbane"),
    ("Milton", -27.4700, 153.0030, "Inner Brisbane"), ("Toowong", -27.4850, 152.9930, "Brisbane West"),
    ("Indooroopilly", -27.4990, 152.9730, "Brisbane West"), ("Kenmore", -27.5070, 152.9380, "Brisbane West"),
    ("The Gap", -27.4430, 152.9480, "Brisbane West"), ("Mitchelton", -27.4170, 152.9780, "Brisbane North"),
    ("Everton Park", -27.4040, 152.9900, "Brisbane North"), ("Chermside", -27.3850, 153.0310, "Brisbane North"),
    ("Aspley", -27.3640, 153.0170, "Brisbane North"), ("Nundah", -27.4010, 153.0590, "Brisbane North"),
    ("Virginia", -27.3780, 153.0620, "Brisbane North"), ("Banyo", -27.3730, 153.0800, "Brisbane North"),
    ("Eagle Farm", -27.4300, 153.0870, "Brisbane East"), ("Hamilton", -27.4380, 153.0640, "Brisbane East"),
    ("Pinkenba", -27.4220, 153.1180, "Brisbane East"), ("Bulimba", -27.4520, 153.0590, "Brisbane East"),
    ("Cannon Hill", -27.4700, 153.0900, "Brisbane East"), ("Murarrie", -27.4610, 153.0990, "Brisbane East"),
    ("Wynnum", -27.4420, 153.1730, "Brisbane East"), ("Carindale", -27.5030, 153.1020, "Brisbane East"),
    ("Capalaba", -27.5420, 153.1960, "Brisbane East"), ("Mount Gravatt", -27.5370, 153.0800, "Brisbane South"),
    ("Sunnybank", -27.5800, 153.0600, "Brisbane South"), ("Acacia Ridge", -27.5860, 153.0260, "Brisbane South"),
    ("Archerfield", -27.5690, 153.0200, "Brisbane South"), ("Salisbury", -27.5530, 153.0310, "Brisbane South"),
    ("Coopers Plains", -27.5650, 153.0400, "Brisbane South"), ("Rocklea", -27.5390, 152.9950, "Brisbane South"),
    ("Oxley", -27.5540, 152.9790, "Brisbane South"), ("Forest Lake", -27.6200, 152.9660, "Brisbane South"),
    ("Underwood", -27.6080, 153.1100, "Logan"), ("Springwood", -27.6130, 153.1290, "Logan"),
    ("Logan Central", -27.6390, 153.1090, "Logan"), ("Browns Plains", -27.6630, 153.0490, "Logan"),
    ("Beenleigh", -27.7150, 153.2010, "Logan"), ("Yatala", -27.7480, 153.2240, "Logan"),
    ("Darra", -27.5670, 152.9530, "Ipswich"), ("Goodna", -27.6100, 152.9000, "Ipswich"),
    ("Redbank", -27.6000, 152.8700, "Ipswich"), ("Springfield", -27.6540, 152.9170, "Ipswich"),
    ("Ipswich", -27.6140, 152.7580, "Ipswich"),
    ("Strathpine", -27.3040, 152.9900, "Moreton Bay"), ("Brendale", -27.3210, 152.9840, "Moreton Bay"),
    ("North Lakes", -27.2350, 153.0200, "Moreton Bay"), ("Redcliffe", -27.2300, 153.1100, "Moreton Bay"),
    ("Morayfield", -27.1080, 152.9490, "Moreton Bay"), ("Caboolture", -27.0850, 152.9510, "Moreton Bay"),
    ("Coomera", -27.8630, 153.3140, "Gold Coast"), ("Nerang", -27.9890, 153.3360, "Gold Coast"),
    ("Southport", -27.9670, 153.4000, "Gold Coast"), ("Surfers Paradise", -28.0020, 153.4300, "Gold Coast"),
    ("Robina", -28.0780, 153.3850, "Gold Coast"), ("Burleigh Heads", -28.0870, 153.4500, "Gold Coast"),
    ("Caloundra", -26.8030, 153.1220, "Sunshine Coast"), ("Maroochydore", -26.6600, 153.1000, "Sunshine Coast"),
    ("Nambour", -26.6270, 152.9590, "Sunshine Coast"), ("Noosa Heads", -26.3940, 153.0900, "Sunshine Coast"),
    ("Toowoomba", -27.5600, 151.9500, "Toowoomba"), ("Toowoomba", -27.5800, 151.9300, "Toowoomba"),
]
KINDS = ["Hardware", "Trade centre", "Pharmacy", "Building supplies", "Garden centre", "Grocer", "Auto parts",
         "Cafe supplies", "Medical clinic", "Homewares", "Pet supplies"]
CARRIER_BY_REGION = {"Inner Brisbane": "Own trucks", "Brisbane West": "Own trucks", "Brisbane North": "Own trucks",
                     "Brisbane East": "Own trucks", "Brisbane South": "Own trucks", "Ipswich": "Own trucks",
                     "Logan": "Carrier A", "Gold Coast": "Carrier A", "Moreton Bay": "Carrier B",
                     "Sunshine Coast": "Carrier B", "Toowoomba": "Carrier B"}
BASE_LATE = {"Own trucks": 0.04, "Carrier A": 0.07, "Carrier B": 0.06, "Same-day courier": 0.03}
LEAD_DAYS = {"Gold Coast": 1, "Sunshine Coast": 2, "Toowoomba": 2}   # working days from order; others next day
OUT_REASONS = [("Carrier delay", 4), ("Run overloaded", 3), ("Customer closed (failed attempt)", 2),
               ("Stock not ready", 2), ("Wrong address", 1)]

# supplier id, name, place, state, lat, lon, transit working days, carrier, usual chance of being late
SUPPLIERS = [
    ("S01", "Supplier A", "Wetherill Park", "NSW", -33.850, 150.900, 2, "Line-haul (road)", 0.08),
    ("S02", "Supplier B", "Dandenong South", "VIC", -38.020, 145.210, 3, "Line-haul (road)", 0.10),
    ("S03", "Supplier C", "Laverton North", "VIC", -37.830, 144.790, 3, "Line-haul (road)", 0.10),
    ("S04", "Supplier D (imports)", "Port of Brisbane", "QLD", -27.380, 153.170, 1, "Container cartage", 0.10),
    ("S05", "Supplier E", "Yatala", "QLD", -27.750, 153.220, 1, "Supplier's own truck", 0.04),
    ("S06", "Supplier F", "Narangba", "QLD", -27.200, 152.960, 1, "Supplier's own truck", 0.05),
    ("S07", "Supplier G", "Toowoomba", "QLD", -27.560, 151.950, 1, "Supplier's own truck", 0.06),
    ("S08", "Supplier H", "Wingfield", "SA", -34.850, 138.570, 4, "Line-haul (road)", 0.10),
    ("S09", "Supplier I", "Kewdale", "WA", -31.980, 115.950, 6, "Line-haul (rail)", 0.12),
    ("S10", "Supplier J", "Beresfield", "NSW", -32.800, 151.650, 2, "Line-haul (road)", 0.07),
    ("S11", "Supplier K", "Townsville", "QLD", -19.260, 146.820, 3, "Line-haul (road)", 0.09),
    ("S12", "Supplier L", "Rocklea", "QLD", -27.540, 152.990, 1, "Supplier's own truck", 0.03),
    ("S13", "Supplier M (air)", "Brisbane Airport", "QLD", -27.390, 153.120, 1, "Air freight", 0.05),
    ("S14", "Supplier N", "Bundamba", "QLD", -27.610, 152.800, 1, "Supplier's own truck", 0.04),
]
IN_REASONS = [("Line-haul delay", 4), ("Supplier stock shortage", 3), ("Booking slot missed", 2), ("Port / customs hold", 1)]
COLOURS = {"On time": "#0E9F6E", "Late": "#b28a92", "Very late": "#8f4a3e", "Overdue": "#8f4a3e", "In transit": "#8e9cab",
           "Due": "#8e9cab"}


def wdays_between(a, b):
    ds, d = [], a
    while d <= b:
        if working(d):
            ds.append(d)
        d += timedelta(days=1)
    return ds


def add_wdays(d, n):
    while n > 0:
        d += timedelta(days=1)
        if working(d):
            n -= 1
    return d


def sub_wdays(d, n):
    while n > 0:
        d -= timedelta(days=1)
        if working(d):
            n -= 1
    return d


def pick(rng, weighted):
    return rng.choices([w[0] for w in weighted], [w[1] for w in weighted])[0]


# ------------------------------------------------------------------ generate

def customers():
    rng = random.Random(66)
    rows = []
    for i, (sub, lat, lon, region) in enumerate(SUBURBS, 1):
        kind = KINDS[(i * 7) % len(KINDS)]
        rows.append({"customer_id": f"C{i:03d}", "name": f"{sub} {kind}" + (" 2" if sub == "Toowoomba" and i % 2 == 0 else ""),
                     "suburb": sub, "region": region, "lat": round(lat + rng.uniform(-0.006, 0.006), 4),
                     "lon": round(lon + rng.uniform(-0.006, 0.006), 4), "carrier": CARRIER_BY_REGION[region],
                     "deliveries_per_week": rng.choice([1, 2, 2, 3, 3, 4, 5])})
    return rows


def deliveries_out(custs):
    rng = random.Random(2026)
    rows, n = [], 0
    for day in wdays_between(START, TODAY):
        for c in custs:
            if rng.random() >= c["deliveries_per_week"] / 5:
                continue
            n += 1
            carrier = "Same-day courier" if c["region"] == "Inner Brisbane" and rng.random() < 0.15 else c["carrier"]
            p_late = BASE_LATE[carrier] + (0.05 if day.weekday() == 0 else 0)
            if carrier == "Carrier A" and c["region"] == "Gold Coast" and day >= date(2026, 9, 7):
                p_late = 0.38           # the Gold Coast run goes wrong
            late = rng.random() < p_late
            days_late = 0
            reason = None
            if late:
                days_late = 1 if rng.random() < 0.72 else rng.choice([2, 2, 3])
                reason = "Carrier delay" if carrier == "Carrier A" and c["region"] == "Gold Coast" and rng.random() < 0.7 else pick(rng, OUT_REASONS)
            delivered = add_wdays(day, days_late) if days_late else day
            hour = rng.randint(7, 16)
            status = "On time" if not days_late else ("Late" if days_late == 1 else "Very late")
            delivered_ts = datetime(delivered.year, delivered.month, delivered.day, hour, rng.choice([0, 10, 20, 30, 40, 50]))
            if delivered > TODAY or (delivered == TODAY and hour >= NOW_HOUR):   # not delivered yet
                status = "In transit" if day == TODAY else "Overdue"
                delivered_ts = None
                days_late = len(wdays_between(day + timedelta(days=1), TODAY)) if day < TODAY else 0
            cartons = rng.randint(2, 40)
            rows.append({"delivery_id": f"D{n:05d}", "order_id": f"SO{40000 + n}", "customer_id": c["customer_id"],
                         "carrier": carrier, "order_date": sub_wdays(day, LEAD_DAYS.get(c["region"], 1)),
                         "promised_date": day, "delivered_at": delivered_ts, "status": status, "days_late": days_late,
                         "late_reason": reason if status in ("Late", "Very late", "Overdue") else None,
                         "cartons": cartons, "order_value": cartons * rng.randint(55, 140),
                         "in_full": 0 if rng.random() < 0.03 else 1})
    return rows


def deliveries_in():
    rng = random.Random(4040)
    rows, n = [], 0
    for day in wdays_between(START, TODAY + timedelta(days=4)):   # includes POs due in the next few days
        for s in SUPPLIERS:
            sid, name, place, state, lat, lon, transit, carrier, p_late = s
            if rng.random() >= (0.55 if state == "QLD" else 0.3):
                continue
            n += 1
            if sid == "S03" and date(2026, 9, 1) <= day <= date(2026, 9, 30):
                p_late = 0.55        # Melbourne line-haul problems
            late = rng.random() < p_late
            delay = (1 if rng.random() < 0.6 else rng.choice([2, 3, 4])) if late else 0
            reason = None
            if late:
                reason = "Line-haul delay" if sid == "S03" else pick(rng, IN_REASONS)
            if sid == "S04" and date(2026, 9, 21) <= day <= date(2026, 9, 25):
                delay, reason = 6, "Port / customs hold"      # container held at the port
            arrived = add_wdays(day, delay) if delay else day
            if day > TODAY:
                status, arrived, dl = "Due", None, 0
            elif arrived > TODAY or (arrived == TODAY and rng.random() < 0.5):
                status = "Overdue" if day < TODAY else "Due"
                dl = len(wdays_between(day + timedelta(days=1), TODAY)) if day < TODAY else 0
                arrived = None
            else:
                dl = delay
                status = "On time" if not delay else ("Late" if delay == 1 else "Very late")
            lines = rng.randint(3, 30)
            rows.append({"receipt_id": f"R{n:05d}", "po_id": f"PO{7000 + n}", "supplier_id": sid, "carrier": carrier,
                         "po_date": sub_wdays(day, transit + rng.randint(2, 6)), "due_date": day, "arrived_date": arrived,
                         "status": status, "days_late": dl, "late_reason": reason if status in ("Late", "Very late", "Overdue") else None,
                         "lines": lines, "lines_short": (rng.randint(1, 3) if rng.random() < 0.06 else 0) if arrived else 0,
                         "po_value": lines * rng.randint(180, 900)})
    return rows


# ------------------------------------------------------------------ periods

def periods():
    last_wd = sub_wdays(TODAY, 1)
    return [
        {"id": "today", "label": "Today", "long": f"Today, {TODAY.strftime('%a %-d %b')} (as at 2pm)", "from": TODAY, "to": TODAY},
        {"id": "last", "label": "Last working day", "long": f"{last_wd.strftime('%a %-d %b')} (last working day)", "from": last_wd, "to": last_wd},
        {"id": "7d", "label": "7 days", "long": f"Last 7 days ({(TODAY - timedelta(days=6)).strftime('%-d %b')} to {TODAY.strftime('%-d %b')})",
         "from": TODAY - timedelta(days=6), "to": TODAY},
        {"id": "30d", "label": "30 days", "long": f"Last 30 days ({(TODAY - timedelta(days=29)).strftime('%-d %b')} to {TODAY.strftime('%-d %b')})",
         "from": TODAY - timedelta(days=29), "to": TODAY},
    ]


def summarise(rows, date_field, place_field, places, group_label):
    """Per period: KPIs, one point per place, and the worst hotspots."""
    out = {}
    for p in periods():
        sel = [r for r in rows if p["from"] <= r[date_field] <= p["to"]]
        done = [r for r in sel if r["status"] in ("On time", "Late", "Very late")]
        ontime = sum(r["status"] == "On time" for r in done)
        pts = {}
        for r in sel:
            k = r[place_field]
            pt = pts.setdefault(k, {"id": k, "n": 0, "on_time": 0, "late": 0, "very_late": 0, "overdue": 0, "open": 0})
            pt["n"] += 1
            pt[{"On time": "on_time", "Late": "late", "Very late": "very_late", "Overdue": "overdue",
                "In transit": "open", "Due": "open"}[r["status"]]] += 1
        groups = {}
        for r in sel:
            g = group_label(r)
            a = groups.setdefault(g, [0, 0])
            a[0] += 1
            a[1] += r["status"] in ("Late", "Very late", "Overdue")
        hot = sorted([[g, a[1], a[0]] for g, a in groups.items() if a[0] >= 6 and a[1]], key=lambda x: (-x[1] / x[2], -x[2]))[:3]
        badge, note = period_status(p["from"], p["to"])
        out[p["id"]] = {
            "status": badge, "status_note": note,
            "kpis": {"total": len(sel), "on_time_pct": round(100 * ontime / len(done), 1) if done else None,
                     "late": sum(r["status"] in ("Late", "Very late") for r in sel),
                     "overdue": sum(r["status"] == "Overdue" for r in sel),
                     "open": sum(r["status"] in ("In transit", "Due") for r in sel),
                     "in_full_pct": round(100 * sum(r.get("in_full", 1 if not r.get("lines_short") else 0) for r in done) / len(done), 1) if done else None},
            "points": [dict(pt, **places[pt["id"]]) for pt in pts.values()],
            "hotspots": [{"label": h[0], "late": h[1], "of": h[2], "pct": round(100 * h[1] / h[2], 1)} for h in hot],
        }
    return out


# ------------------------------------------------------------------- build

def build():
    custs = customers()
    out_rows = deliveries_out(custs)
    in_rows = deliveries_in()
    check(out_rows, in_rows)
    cplace = {c["customer_id"]: {"name": c["name"], "sub": c["suburb"], "region": c["region"], "lat": c["lat"], "lon": c["lon"]} for c in custs}
    splace = {s[0]: {"name": s[1], "sub": f"{s[2]} {s[3]}", "region": s[3], "lat": s[4], "lon": s[5]} for s in SUPPLIERS}
    region = {c["customer_id"]: c["region"] for c in custs}
    sname = {s[0]: s[1] for s in SUPPLIERS}
    payload = {
        "as_at": f"{TODAY.isoformat()}T{NOW_HOUR:02d}:00", "dc": DC, "colours": COLOURS,
        "periods": [{"id": p["id"], "label": p["label"], "long": p["long"]} for p in periods()],
        "out": summarise(out_rows, "promised_date", "customer_id", cplace, lambda r: f"{r['carrier']} · {region[r['customer_id']]}"),
        "in": summarise(in_rows, "due_date", "supplier_id", splace, lambda r: f"{sname[r['supplier_id']]} · {r['carrier']}"),
        "exports": {"xlsx": "media/exports/deliveries-sample.xlsx", "pdf": "media/exports/deliveries-sample.pdf", "pptx": "media/exports/deliveries-sample.pptx"},
    }
    return dict(customers=custs, out=out_rows, inn=in_rows, payload=payload)


def check(out_rows, in_rows):
    problems = []
    for r in out_rows:
        if r["status"] in ("On time", "Late", "Very late") and r["delivered_at"] is None: problems.append(f"{r['delivery_id']} no delivery time")
        if r["status"] == "On time" and r["delivered_at"].date() != r["promised_date"]: problems.append(f"{r['delivery_id']} on time but late")
        if r["status"] in ("Late", "Very late") and r["delivered_at"].date() <= r["promised_date"]: problems.append(f"{r['delivery_id']} late but on time")
        if r["promised_date"] > TODAY or r["order_date"] >= r["promised_date"]: problems.append(f"{r['delivery_id']} dates")
        if not working(r["promised_date"]) or r["promised_date"] in QLD_HOLIDAYS: problems.append(f"{r['delivery_id']} holiday")
    for r in in_rows:
        if r["status"] in ("On time", "Late", "Very late") and r["arrived_date"] is None: problems.append(f"{r['receipt_id']} no arrival")
        if r["status"] == "Due" and r["due_date"] < TODAY: problems.append(f"{r['receipt_id']} due in the past")
    if problems:
        raise SystemExit("Deliveries checks FAILED:\n  " + "\n  ".join(problems[:20]))


def tables(d):
    V = lambda n: f"VARCHAR({n})"
    T = {}
    def tbl(name, desc, cols, rows):
        T[name] = (desc, [c[0] for c in cols], rows, cols)
    tbl("dim_customer", "Deliveries map: Sample Distribution's customers across south-east Queensland, with map coordinates and usual carrier.",
        [("customer_id", V(6) + " PRIMARY KEY"), ("customer_name", V(60)), ("suburb", V(40)), ("region", V(30)),
         ("latitude", "DECIMAL(9,4)"), ("longitude", "DECIMAL(9,4)"), ("usual_carrier", V(30)), ("deliveries_per_week", "INT")],
        [[c["customer_id"], c["name"], c["suburb"], c["region"], c["lat"], c["lon"], c["carrier"], c["deliveries_per_week"]] for c in d["customers"]])
    tbl("dim_supplier", "Deliveries map: suppliers delivering into the Wacol DC, with map coordinates and normal transit time.",
        [("supplier_id", V(6) + " PRIMARY KEY"), ("supplier_name", V(40)), ("place", V(40)), ("state", V(3)),
         ("latitude", "DECIMAL(9,4)"), ("longitude", "DECIMAL(9,4)"), ("transit_working_days", "INT"), ("usual_carrier", V(30))],
        [[s[0], s[1], s[2], s[3], s[4], s[5], s[6], s[7]] for s in SUPPLIERS])
    tbl("fact_delivery_out", "Deliveries map: every delivery out of the Wacol DC, 3 Aug to 6 Oct 2026 (as at 2pm on 6 Oct). status: On time / Late (1 working day) / Very late (2+) / Overdue (not delivered, past promised date) / In transit (due today). days_late in working days.",
        [("delivery_id", V(8) + " PRIMARY KEY"), ("order_id", V(10)), ("customer_id", V(6) + " REFERENCES dim_customer(customer_id)"),
         ("carrier", V(30)), ("order_date_key", "INT REFERENCES dim_date(date_key)"), ("promised_date_key", "INT REFERENCES dim_date(date_key)"),
         ("delivered_date_key", "INT REFERENCES dim_date(date_key)"), ("delivered_time", V(5)), ("status", V(12)), ("days_late", "INT"), ("late_reason", V(40)),
         ("cartons", "INT"), ("order_value", "INT"), ("in_full", "INT")],
        [[r["delivery_id"], r["order_id"], r["customer_id"], r["carrier"], key(r["order_date"]), key(r["promised_date"]),
          key(r["delivered_at"].date()) if r["delivered_at"] else None, r["delivered_at"].strftime("%H:%M") if r["delivered_at"] else None, r["status"], r["days_late"], r["late_reason"],
          r["cartons"], r["order_value"], r["in_full"]] for r in d["out"]])
    tbl("fact_delivery_in", "Deliveries map: every purchase-order delivery into the Wacol DC, due 3 Aug to 9 Oct 2026 (as at 6 Oct). status: On time / Late / Very late / Overdue / Due (not due yet). lines_short = lines not delivered in full.",
        [("receipt_id", V(8) + " PRIMARY KEY"), ("po_id", V(10)), ("supplier_id", V(6) + " REFERENCES dim_supplier(supplier_id)"),
         ("carrier", V(30)), ("po_date_key", "INT REFERENCES dim_date(date_key)"), ("due_date_key", "INT REFERENCES dim_date(date_key)"),
         ("arrived_date_key", "INT REFERENCES dim_date(date_key)"), ("status", V(12)), ("days_late", "INT"), ("late_reason", V(40)),
         ("lines", "INT"), ("lines_short", "INT"), ("po_value", "INT")],
        [[r["receipt_id"], r["po_id"], r["supplier_id"], r["carrier"], key(r["po_date"]), key(r["due_date"]),
          key(r["arrived_date"]) if r["arrived_date"] else None, r["status"], r["days_late"], r["late_reason"],
          r["lines"], r["lines_short"], r["po_value"]] for r in d["inn"]])
    return T


def write_js(d):
    js = ("/* deliveries-data.js — GENERATED by tools/sample_data.py from tools/deliveries.py. Don't edit by hand.\n"
          "   Invented sample data: deliveries in and out of a Brisbane distribution centre, aggregated for the map. */\n"
          "window.FOURTH_SHEET_DELIVERIES = " + json.dumps(d["payload"], ensure_ascii=False, separators=(",", ":")) + ";\n")
    (REPO / "data" / "deliveries-data.js").write_text(js)


# ------------------------------------------------------------------ Excel

def write_xlsx(d):
    """Every delivery as data; the summary and the late tables are COUNTIFS formulas over those rows."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    GREEN, INK, MUTED = "2F7A5D", "25342A", "5F6F63"
    F = ek.XL_FMT
    THIN = Side(style="thin", color="DCE8DC")
    business = "Sample Distribution Pty Ltd"
    filename = "deliveries-sample.xlsx"
    retrieved = ds.as_at_text()
    sub = "Sample Distribution · Wacol DC · SAMPLE DATA (invented) · The Fourth Sheet"
    status = (f"Data status (data retrieved {retrieved}): today's deliveries are incomplete; 1–5 October is provisional "
              f"(October locks {ds.lock_date('2026-10').strftime('%-d %b')}); September and earlier are locked.")
    wb = Workbook()
    pages = {}

    def title(ws, t):
        for cell, val, font in (("A1", t, Font(bold=True, size=14, color=INK)), ("A2", sub, Font(italic=True, size=9, color=MUTED)),
                                ("A3", status, Font(italic=True, size=9, color=MUTED))):
            ws[cell] = val
            ws[cell].font = font

    def head(ws, r, cols, widths):
        for c, h in enumerate(cols, 1):
            cell = ws.cell(r, c, h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=GREEN)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(c)].width = widths.get(h, 13)

    def data_sheet(name, t, cols, kinds, rows, widths, group=None):
        ws = wb.create_sheet(name)
        title(ws, t)
        hr = 5
        head(ws, hr, cols, widths)
        starts = []
        for i, row_ in enumerate(rows):
            r = hr + 1 + i
            if group and i and group[i] != group[i - 1]:
                starts.append(r)
            for c, (val, kind) in enumerate(zip(row_, kinds), 1):
                if isinstance(val, str) and val.startswith("="):
                    val = val.format(r=r)
                cell = ws.cell(r, c, val)
                cell.number_format = F[kind]
        last = hr + len(rows)
        ws.freeze_panes = f"B{hr + 1}"
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(cols))}{last}"
        pages[name] = (hr, starts, last)
        return ws, last

    # ---- data sheets
    custs = sorted(d["customers"], key=lambda c: (c["region"], c["name"]))
    data_sheet("Customers", "Customers", ["Customer ID", "Customer", "Suburb", "Region", "Latitude", "Longitude", "Usual carrier", "Deliveries a week"],
               ["text", "text", "text", "text", "text", "text", "text", "int"],
               [[c["customer_id"], c["name"], c["suburb"], c["region"], c["lat"], c["lon"], c["carrier"], c["deliveries_per_week"]] for c in custs],
               {"Customer": 30, "Suburb": 18, "Region": 16, "Usual carrier": 18}, group=[c["region"] for c in custs])
    data_sheet("Suppliers", "Suppliers", ["Supplier ID", "Supplier", "Place", "State", "Latitude", "Longitude", "Transit (working days)", "Usual carrier"],
               ["text", "text", "text", "text", "text", "text", "int", "text"],
               [[sp[0], sp[1], sp[2], sp[3], sp[4], sp[5], sp[6], sp[7]] for sp in SUPPLIERS],
               {"Supplier": 22, "Place": 18, "Usual carrier": 20})
    out = sorted(d["out"], key=lambda r: (r["promised_date"], r["delivery_id"]))
    look_c = lambda col: f"=INDEX(Customers!${col}:${col},MATCH(C{{r}},Customers!$A:$A,0))"
    ws_out, last_out = data_sheet(
        "Deliveries out", "Every delivery out, 3 Aug to 6 Oct 2026",
        ["Delivery", "Order", "Customer ID", "Customer", "Region", "Carrier", "Order date", "Promised", "Delivered", "Time", "Status",
         "Working days late", "Late reason", "Cartons", "Order value", "In full", "Completed"],
        ["text", "text", "text", "text", "text", "text", "date", "date", "date", "text", "text", "int", "text", "int", "money", "int", "int"],
        [[r["delivery_id"], r["order_id"], r["customer_id"], look_c("B"), look_c("D"), r["carrier"], r["order_date"], r["promised_date"],
          r["delivered_at"].date() if r["delivered_at"] else None, r["delivered_at"].strftime("%H:%M") if r["delivered_at"] else None,
          r["status"], r["days_late"], r["late_reason"], r["cartons"], r["order_value"], r["in_full"],
          '=IF(OR(K{r}="On time",K{r}="Late",K{r}="Very late"),1,0)'] for r in out],
        {"Customer": 28, "Region": 16, "Carrier": 17, "Late reason": 28, "Status": 11}, group=[r["promised_date"].strftime("%Y-%m") for r in out])
    inn = sorted(d["inn"], key=lambda r: (r["due_date"], r["receipt_id"]))
    ws_in, last_in = data_sheet(
        "Deliveries in", "Every purchase-order delivery in, due 3 Aug to 9 Oct 2026",
        ["Receipt", "PO", "Supplier ID", "Supplier", "Carrier", "PO date", "Due", "Arrived", "Status", "Working days late", "Late reason",
         "Lines", "Lines short", "PO value", "Completed", "In full"],
        ["text", "text", "text", "text", "text", "date", "date", "date", "text", "int", "text", "int", "int", "money", "int", "int"],
        [[r["receipt_id"], r["po_id"], r["supplier_id"], "=INDEX(Suppliers!$B:$B,MATCH(C{r},Suppliers!$A:$A,0))", r["carrier"], r["po_date"],
          r["due_date"], r["arrived_date"], r["status"], r["days_late"], r["late_reason"], r["lines"], r["lines_short"], r["po_value"],
          '=IF(OR(I{r}="On time",I{r}="Late",I{r}="Very late"),1,0)', "=IF(AND(O{r}=1,M{r}=0),1,0)"] for r in inn],
        {"Supplier": 22, "Carrier": 20, "Late reason": 26, "Status": 11}, group=[r["due_date"].strftime("%Y-%m") for r in inn])
    sand, brick = PatternFill("solid", fgColor="F3EADF"), PatternFill("solid", fgColor="EFD9D5")
    for ws_, col, last in ((ws_out, "K", last_out), (ws_in, "I", last_in)):
        rng_ = f"{col}6:{col}{last}"
        ws_.conditional_formatting.add(rng_, CellIsRule(operator="equal", formula=['"Late"'], fill=sand))
        for st_ in ("Very late", "Overdue"):
            ws_.conditional_formatting.add(rng_, CellIsRule(operator="equal", formula=[f'"{st_}"'], fill=brick))

    # ---- summary: periods are inputs, everything else is a formula over the data sheets
    ws = wb.active
    ws.title = "Summary"
    title(ws, "Deliveries in and out: summary")
    r, starts = 5, []
    cols = ["Period", "From", "To", "Due in period", "On time", "Late", "Very late", "Overdue now", "Not due yet / on its way", "On time %", "In full %"]
    widths = {"Period": 30, "Not due yet / on its way": 15}
    for side, label, sh, dcol, scol, fcol, ccol in (("out", "Deliveries out, to customers", "'Deliveries out'", "H", "K", "P", "Q"),
                                                     ("in", "Deliveries in, from suppliers", "'Deliveries in'", "G", "I", "P", "O")):
        starts.append(r)
        ws.cell(r, 1, label).font = Font(bold=True, color=GREEN, size=12)
        r += 1
        head(ws, r, cols, widths)
        for p in periods():
            r += 1
            ws.cell(r, 1, p["long"])
            ws.cell(r, 2, p["from"]).number_format = F["date"]
            ws.cell(r, 3, p["to"]).number_format = F["date"]
            rng = f'{sh}!${dcol}:${dcol},">="&$B{r},{sh}!${dcol}:${dcol},"<="&$C{r}'
            ws.cell(r, 4, f"=COUNTIFS({rng})")
            for c, st_ in ((5, "On time"), (6, "Late"), (7, "Very late"), (8, "Overdue")):
                ws.cell(r, c, f'=COUNTIFS({rng},{sh}!${scol}:${scol},"{st_}")')
            open_st = "In transit" if side == "out" else "Due"
            ws.cell(r, 9, f'=COUNTIFS({rng},{sh}!${scol}:${scol},"{open_st}")')
            ws.cell(r, 10, f'=IF(SUM(E{r}:G{r})=0,"",E{r}/SUM(E{r}:G{r}))').number_format = F["pct"]
            ws.cell(r, 11, f'=IF(SUM(E{r}:G{r})=0,"",SUMIFS({sh}!${fcol}:${fcol},{rng},{sh}!${ccol}:${ccol},1)/SUM(E{r}:G{r}))').number_format = F["pct"]
            for c in range(4, 10):
                ws.cell(r, c).number_format = F["int"]
        r += 3
    # where it's going wrong, last 30 days (rows listed worst first; every number is a formula)
    p30 = periods()[3]
    for side, label, sh, dcol, scol in (("out", "Late by carrier and region, last 30 days", "'Deliveries out'", "H", "K"),
                                        ("in", "Late by supplier, last 30 days", "'Deliveries in'", "G", "I")):
        starts.append(r)
        ws.cell(r, 1, label).font = Font(bold=True, color=GREEN, size=12)
        ws.cell(r + 1, 1, "From").font = Font(color=MUTED)
        ws.cell(r + 1, 2, p30["from"]).number_format = F["date"]
        ws.cell(r + 1, 3, p30["to"]).number_format = F["date"]
        fr = r + 1
        r += 2
        head(ws, r, ["Carrier · region" if side == "out" else "Supplier", "Carrier" if side == "out" else "Supplier ID",
                     "Region" if side == "out" else "", "Due in period", "Late, very late or overdue", "% late"], widths)
        rows_src = d["out"] if side == "out" else d["inn"]
        sel = [x for x in rows_src if p30["from"] <= (x["promised_date"] if side == "out" else x["due_date"]) <= p30["to"]]
        reg = {c["customer_id"]: c["region"] for c in d["customers"]}
        if side == "out":
            groups = sorted({(x["carrier"], reg[x["customer_id"]]) for x in sel})
        else:
            groups = sorted({(x["supplier_id"], "") for x in sel})
        def rate(g):
            m_ = [x for x in sel if ((x["carrier"], reg[x["customer_id"]]) if side == "out" else (x["supplier_id"], "")) == g]
            return -sum(x["status"] in ("Late", "Very late", "Overdue") for x in m_) / len(m_)
        sname = {sp[0]: sp[1] for sp in SUPPLIERS}
        for g in sorted(groups, key=rate):
            r += 1
            base = f'{sh}!${dcol}:${dcol},">="&$B${fr},{sh}!${dcol}:${dcol},"<="&$C${fr}'
            if side == "out":
                ws.cell(r, 1, f"{g[0]} · {g[1]}")
                ws.cell(r, 2, g[0]); ws.cell(r, 3, g[1])
                crit = f'{base},{sh}!$F:$F,$B{r},{sh}!$E:$E,$C{r}'
            else:
                ws.cell(r, 1, sname[g[0]]); ws.cell(r, 2, g[0])
                crit = f'{base},{sh}!$C:$C,$B{r}'
            ws.cell(r, 4, f"=COUNTIFS({crit})").number_format = F["int"]
            ws.cell(r, 5, "=" + "+".join(f'COUNTIFS({crit},{sh}!${scol}:${scol},"{st_}")' for st_ in ("Late", "Very late", "Overdue"))).number_format = F["int"]
            ws.cell(r, 6, f'=IF(D{r}=0,"",E{r}/D{r})').number_format = F["pct"]
        r += 3
    for c, wdt in zip("ABCDEFGHIJK", (34, 13, 16, 12, 12, 12, 12, 12, 15, 12, 12)):
        ws.column_dimensions[c].width = wdt
    pages["Summary"] = (None, starts, r)

    review = []
    for ws_ in wb.worksheets:
        hr, st, last = pages[ws_.title]
        ek.xl_print(ws_, business, filename, retrieved, landscape=True, header_row=hr)
        review.append(ek.xl_review(ws_, ek.xl_breaks(ws_, st, last, hr), hr))
    w = wb["Summary"]
    r = w.max_row + 2
    w.cell(r, 1, "Print check (done before this file was saved)").font = Font(bold=True, color=GREEN)
    for line in review:
        r += 1
        w.cell(r, 1, "✓ " + line)
    print(f"  {filename}: " + " | ".join(review))
    EXPORTS.mkdir(parents=True, exist_ok=True)
    ek.save_if_changed(wb.save, EXPORTS / filename)


# -------------------------------------------------------------------- PDF

def html_report(d):
    p = d["payload"]
    esc = lambda s: str(s).replace("&", "&amp;").replace("<", "&lt;")

    def side(s, heading):
        k = p[s]["30d"]["kpis"]
        hot = "".join(f"<li><b>{esc(h['label'])}</b>: {h['late']} of {h['of']} late ({h['pct']:.1f}%)</li>" for h in p[s]["30d"]["hotspots"])
        return f"""<section class=pg><h2>{heading} · last 30 days</h2>
<div class=kpis><div><span>Deliveries</span><b>{k['total']:,}</b></div><div><span>On time</span><b>{k['on_time_pct']:.1f}%</b></div>
<div><span>Late</span><b>{k['late']}</b></div><div><span>Overdue now</span><b>{k['overdue']}</b></div><div><span>In full</span><b>{k['in_full_pct']:.1f}%</b></div></div>
<div class=row><div id=map-{s} class=map></div><div class=side><h3>Where it's going wrong</h3><ul>{hot}</ul>
<p class=legend><i style="background:#0E9F6E"></i>95%+ on time <i style="background:#b28a92"></i>80–95% <i style="background:#8f4a3e"></i>under 80% or overdue</p>
<p class=note>Dot size = number of deliveries. Every delivery is in the Excel download.</p></div></div></section>"""

    return f"""<!doctype html><html><head><meta charset=utf-8>
<link rel=stylesheet href="https://cdnjs.cloudflare.com/ajax/libs/leaflet/1.9.4/leaflet.min.css">
<script src="https://cdnjs.cloudflare.com/ajax/libs/leaflet/1.9.4/leaflet.min.js"></script>
<link rel=stylesheet href="https://fonts.googleapis.com/css2?family=Inter:wght@400;600;700;800&display=swap">
<style>@page {{ size: A4 landscape; margin: 12mm; }} body {{ font-family: Inter, Arial, sans-serif; color: #25342a; font-size: 10pt; margin: 0; }}
header {{ display: flex; align-items: center; gap: 10px; border-bottom: 3px solid #2f7a5d; padding-bottom: 6px; margin-bottom: 8px; }}
.mark {{ min-width: 28px; height: 28px; padding: 0 4px; border-radius: 7px; background: #2f7a5d; color: #fff; font-weight: 800; display: flex; align-items: center; justify-content: center; }}
sup {{ font-size: .55em; }} header small {{ margin-left: auto; color: #5f6f63; }}
.sample {{ display: inline-block; background: #f6f1e7; border: 1px solid #b28a92; color: #6b5532; border-radius: 4px; padding: 1px 7px; font-size: 8pt; font-weight: 700; }}
h1 {{ font-size: 15pt; margin: 4px 0; }} h2 {{ font-size: 12pt; color: #2f7a5d; margin: 6px 0; }} h3 {{ font-size: 10.5pt; margin: 0 0 4px; }}
.pg + .pg {{ page-break-before: always; }} .kpis {{ display: grid; grid-template-columns: repeat(5, 1fr); gap: 6px; margin-bottom: 8px; }}
.kpis div {{ border: 1px solid #dce8dc; border-radius: 6px; padding: 5px 8px; }} .kpis span {{ display: block; font-size: 8pt; color: #5f6f63; }} .kpis b {{ font-size: 14pt; color: #2f7a5d; }}
.row {{ display: grid; grid-template-columns: 2fr 1fr; gap: 12px; }} .map {{ height: 96mm; border: 1px solid #dce8dc; border-radius: 6px; }}
.side ul {{ padding-left: 16px; }} .side li {{ margin: 4px 0; }} .legend i {{ display: inline-block; width: 9px; height: 9px; border-radius: 50%; margin: 0 3px 0 8px; }}
.note {{ color: #5f6f63; font-size: 8.5pt; }} .leaflet-control-attribution {{ font-size: 7px; }}
""" + ek.STATUS_CSS + f"""</style></head><body>
<header><div class=mark>4<sup>th</sup></div><b>The 4<sup>th</sup> Sheet</b><small>Deliveries in and out · as at {TODAY.strftime('%-d %b %Y')}, 2pm</small></header>
<span class=sample>SAMPLE DATA · invented figures for demonstration</span>
<h1>Sample Distribution: Wacol DC</h1>
{ek.status_box_html([("Today (6 Oct)", "Incomplete", f"As at {ds.as_at_text()}: deliveries still on the road are shown as on their way."),
                     ("1–5 October", "Provisional", f"October is not locked until {ds.lock_date('2026-10').strftime('%-d %b')}: late proofs of delivery can still change it."),
                     ("September 2026 and earlier", "Locked", "Closed. These numbers won't change.")])}
{side('out', 'Deliveries out, to customers')}
{side('in', 'Deliveries in, from suppliers')}
<script>
const D = {json.dumps(p)};
function band(pt) {{ const done = pt.on_time + pt.late + pt.very_late; if (pt.overdue) return '#8f4a3e'; if (!done) return '#8e9cab';
  const r = pt.on_time / done; return r >= 0.95 ? '#0E9F6E' : r >= 0.8 ? '#b28a92' : '#8f4a3e'; }}
window.tilesLoaded = 0;
for (const s of ['out', 'in']) {{
  const m = L.map('map-' + s, {{ zoomControl: false, attributionControl: true }});
  const t = L.tileLayer('https://tile.openstreetmap.org/{{z}}/{{x}}/{{y}}.png', {{ attribution: '© OpenStreetMap contributors' }}).addTo(m);
  t.on('load', () => window.tilesLoaded++);
  const pts = D[s]['30d'].points, ll = pts.map(p => [p.lat, p.lon]).concat([[D.dc.lat, D.dc.lon]]);
  pts.forEach(p => L.circleMarker([p.lat, p.lon], {{ radius: 3 + Math.sqrt(p.n) * 1.3, color: '#fff', weight: 1, fillColor: band(p), fillOpacity: .9 }}).addTo(m));
  L.circleMarker([D.dc.lat, D.dc.lon], {{ radius: 7, color: '#25342a', weight: 2, fillColor: '#fff', fillOpacity: 1 }}).bindTooltip('Wacol DC', {{ permanent: true, direction: 'right' }}).addTo(m);
  m.fitBounds(ll, {{ padding: [14, 14] }});
}}
</script></body></html>"""


def write_pdf(d):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        try:
            browser = pw.chromium.launch(channel="chrome")
        except Exception:
            browser = pw.chromium.launch()
        page = browser.new_page(viewport={"width": 1123, "height": 794})
        # OpenStreetMap's tile policy asks for an identifying referer
        page.set_extra_http_headers({"Referer": "https://nnbooth.github.io/thefourthsheet/"})
        page.set_content(html_report(d), wait_until="networkidle")
        page.wait_for_function("window.tilesLoaded >= 2", timeout=30000)
        page.evaluate("document.fonts.ready")
        for side in ("out", "in"):      # map pictures for the PowerPoint version
            page.locator(f"#map-{side}").screenshot(path=str(EXPORTS / f".map-{side}.png"))
        ek.pdf_if_changed(page, EXPORTS / "deliveries-sample.pdf",
                 **ek.pdf_options("Sample Distribution Pty Ltd", "Deliveries in and out · Sample data", "deliveries-sample.pdf", ds.as_at_text(), landscape=True))
        browser.close()


def write_pptx(d):
    """PowerPoint version of the deliveries PDF (map pictures, KPIs, where it's going wrong)."""
    import pptkit
    p = d["payload"]
    deck = pptkit.Deck("Sample Distribution Pty Ltd", "Deliveries in and out", "deliveries-sample.pptx", ds.as_at_text())
    deck.title_slide("Every delivery out to a customer and in from a supplier, for an invented distribution centre at Wacol, Brisbane.",
                     [f"Today (6 Oct): incomplete, as at {ds.as_at_text()}.", f"1-5 October: provisional (October locks {ds.lock_date('2026-10').strftime('%-d %b')}).",
                      "September and earlier: locked."])
    for side, heading in (("out", "Deliveries out, to customers"), ("in", "Deliveries in, from suppliers")):
        k = p[side]["30d"]["kpis"]
        deck.kpi_slide(f"{heading} · last 30 days", [("Deliveries", f"{k['total']:,}", ""), ("On time", f"{k['on_time_pct']:.1f}%", ""),
                                                     ("Late", str(k["late"]), ""), ("Overdue now", str(k["overdue"]), ""), ("In full", f"{k['in_full_pct']:.1f}%", "")])
        png = EXPORTS / f".map-{side}.png"
        if png.exists():
            deck.image_slide(f"{heading} · map", png, "Green 95%+ on time · rose 80-95% · brick under 80% or overdue · grey-blue on its way. Dot size = number of deliveries.")
        deck.table_slides(f"{heading} · where it's going wrong", ["", "Late", "Of", "% late"],
                          [[h["label"], str(h["late"]), str(h["of"]), f"{h['pct']:.1f}%"] for h in p[side]["30d"]["hotspots"]])
    ek.save_if_changed(deck.save, EXPORTS / "deliveries-sample.pptx")
    for side in ("out", "in"):
        (EXPORTS / f".map-{side}.png").unlink(missing_ok=True)


def publish():
    d = build()
    write_js(d)
    write_xlsx(d)
    write_pdf(d)
    write_pptx(d)
    return d


if __name__ == "__main__":
    d = build()
    for s in ("out", "in"):
        for pid, v in d["payload"][s].items():
            print(s, pid, v["kpis"], [h["label"] + f" {h['pct']}%" for h in v["hotspots"]])
    print(len(d["out"]), "out rows,", len(d["inn"]), "in rows")
