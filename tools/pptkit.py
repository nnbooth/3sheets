#!/usr/bin/env python3
"""
pptkit.py — PowerPoint versions of every PDF export, with one shared look.

16:9 slides built on the template's own slide master and layouts (Title Slide, Title Only), with the theme's fonts
set to Roboto and its colours to the brand, so a client can restyle a whole deck from the master.
Every slide has a real title placeholder, speaker notes (what the slide says), and a footer: file name, date
retrieved, slide X of Y.

Charts are native PowerPoint charts (editable): horizontal bars for categories (green on target, grey below,
red losing money or off budget), area charts for months or days. Targets and break-even are drawn as lines on
the chart (the plot area is laid out exactly, so the lines sit at the right value). Negative numbers are red,
in brackets. Tables split across slides; text columns left-aligned, numbers right-aligned.

combo_slide() puts one idea on one slide: headline numbers across the top, the chart, then "What it shows" and
"What you'd do about it" underneath.
"""

import math
import re

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

GREEN, GREY, INK, MUTED, BRAND, NEG, LINE, WARN = "0E9F6E", "7D8790", "25342A", "5F6F63", "2F7A5D", "B42318", "DCE8DC", "8A6D3B"
FMT = {"pct_var": '0.0"% over";0.0"% under";"on budget"', "money_var": '$#,##0" over";$#,##0" under";"on budget"', "money0": '$#,##0;[Red]($#,##0);"-"',
       "pct1": '0.0"%";[Red](0.0"%")', "cents_int": '0"¢";[Red](0"¢")', "int": '#,##0;[Red](#,##0)'}
W, H = Inches(13.333), Inches(7.5)
ROWS_PER_SLIDE = 14
FONT = "Roboto"         # like the site and the business cards (exportkit.XL_FONT for Excel)
L, R_ = Inches(0.5), Inches(0.5)
CONTENT_W = W - L - R_


def ek_generated():
    import exportkit
    return exportkit.generated_text()


def rgb(h):
    return RGBColor.from_string(h)


def is_neg(t):
    t = str(t).strip()
    return (t.startswith("(") or t.startswith("-")) and any(c.isdigit() for c in t)


def lines_needed(text, size_pt, width_emu):
    """Rough line count for wrapped text: Roboto averages about half its size per character."""
    per_line = max(10, int((width_emu / 12700) / (size_pt * 0.5)))
    return sum(max(1, math.ceil(len(part) / per_line)) for part in str(text).split("\n"))


def looks_numeric(v):
    t = str(v).strip().replace(",", "").replace("$", "").replace("%", "").replace("¢", "").replace("(", "").replace(")", "").replace(" h", "").replace("months", "").strip()
    return t in ("", "-", "–") or bool(re.fullmatch(r"[+-]?\d+(\.\d+)?", t))


class Deck:
    def __init__(self, business, title, filename, retrieved):
        self.p = Presentation()
        self.p.slide_width, self.p.slide_height = W, H
        self.business, self.title, self.filename, self.retrieved = business, title, filename, retrieved
        self.layout_title = self.p.slide_layouts[0]       # Title Slide
        self.layout_body = self.p.slide_layouts[5]        # Title Only: a real title placeholder, the rest is ours
        self._theme()

    def _theme(self):
        """Theme fonts (headings and body) = Roboto; theme colours = the brand. Restyle the master, restyle the deck."""
        master = self.p.slide_masters[0]
        theme = master.part.part_related_by(RT.THEME)
        xml = theme.blob.decode("utf-8")
        xml = re.sub(r'(<a:majorFont>\s*<a:latin typeface=")[^"]*', r"\g<1>" + FONT, xml)
        xml = re.sub(r'(<a:minorFont>\s*<a:latin typeface=")[^"]*', r"\g<1>" + FONT, xml)
        for slot, colour in (("dk2", INK), ("lt2", "F4F7F2"), ("accent1", BRAND), ("accent2", GREEN), ("accent3", GREY), ("accent4", WARN),
                             ("accent5", NEG), ("accent6", "8E9CAB")):
            xml = re.sub(rf"(<a:{slot}>\s*)<a:(?:srgbClr|sysClr)[^>]*/>", rf'\g<1><a:srgbClr val="{colour}"/>', xml)
        theme._blob = xml.encode("utf-8")

    # ---------------------------------------------------------------- basics
    def text(self, slide, x, y, w, h, txt, size=12, bold=False, colour=INK, align=None, anchor=None):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.05)
        if anchor:
            tf.vertical_anchor = anchor
        for k, line in enumerate(str(txt).split("\n")):
            para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            run = para.add_run()
            run.text = line
            run.font.size, run.font.bold, run.font.name = Pt(size), bold, FONT
            run.font.color.rgb = rgb(NEG if is_neg(line) and size >= 18 else colour)
            if align:
                para.alignment = align
        return tb

    def notes(self, slide, text):
        if text:
            slide.notes_slide.notes_text_frame.text = text

    def slide(self, heading, sub=None, notes=None):
        """A content slide on the Title Only layout. Returns (slide, top of the free area). The title gets as many
        lines as it needs (shrinking to 20 pt when long) and the subtitle goes below it, never on top of it."""
        s = self.p.slides.add_slide(self.layout_body)
        bar = s.shapes.add_shape(1, 0, 0, W, Inches(0.08))
        bar.fill.solid(); bar.fill.fore_color.rgb = rgb(BRAND); bar.line.fill.background()
        self.text(s, L, Inches(0.16), Inches(9), Inches(0.28), f"{self.business} · {self.title}", 10, colour=MUTED)
        size = 24 if len(heading) <= 70 else 20
        n = min(3, lines_needed(heading, size, CONTENT_W))
        th = Emu(int(Pt(size).emu * 1.2 * n + Inches(0.12)))
        t = s.shapes.title
        t.left, t.top, t.width, t.height = L, Inches(0.42), CONTENT_W, th
        tf = t.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.text = heading
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT
        for r in para.runs:
            r.font.size, r.font.bold, r.font.name = Pt(size), True, FONT
            r.font.color.rgb = rgb(INK)
        y = Inches(0.42) + th + Inches(0.06)
        if sub:
            sh = Emu(int(Pt(12).emu * 1.25 * lines_needed(sub, 12, CONTENT_W) + Inches(0.08)))
            self.text(s, L, y, CONTENT_W, sh, sub, 12, colour=MUTED)
            y += sh + Inches(0.06)
        self.notes(s, notes or (heading + (". " + sub if sub else "")))
        return s, y

    def footers(self):
        n = len(self.p.slides)
        for i, s in enumerate(self.p.slides, 1):
            y = H - Inches(0.4)
            self.text(s, L, y, Inches(3.2), Inches(0.3), self.filename, 9, colour=MUTED)
            self.text(s, Inches(3.4), y, Inches(6.5), Inches(0.3), f"Data retrieved {self.retrieved} · {ek_generated()} · Sample data · Made with Python (python-pptx)", 9, colour=MUTED, align=PP_ALIGN.CENTER)
            self.text(s, Inches(9.3), y, Inches(3.5), Inches(0.3), f"Slide {i} of {n}", 9, colour=MUTED, align=PP_ALIGN.RIGHT)

    def save(self, path):
        self.footers()
        self.fonts()
        self.p.save(str(path))

    def fonts(self):
        """Roboto everywhere: every run of text, every table cell and every chart (titles, axes, labels, legends)."""
        def runs(tf):
            for para in tf.paragraphs:
                para.font.name = FONT
                for r in para.runs:
                    r.font.name = FONT
        for slide in self.p.slides:
            for sh in slide.shapes:
                if sh.has_text_frame:
                    runs(sh.text_frame)
                if getattr(sh, "has_table", False) and sh.has_table:
                    for row in sh.table.rows:
                        for cell in row.cells:
                            runs(cell.text_frame)
                if getattr(sh, "has_chart", False) and sh.has_chart:
                    sh.chart.font.name = FONT

    # ---------------------------------------------------------------- slides
    def title_slide(self, intro, status_lines):
        s = self.p.slides.add_slide(self.layout_title)
        for ph in list(s.placeholders):
            if ph.placeholder_format.idx != 0:
                ph._element.getparent().remove(ph._element)          # the subtitle placeholder: ours is laid out below
        box = s.shapes.add_shape(1, L, Inches(0.6), Inches(0.7), Inches(0.7))
        box.fill.solid(); box.fill.fore_color.rgb = rgb(BRAND); box.line.fill.background()
        box.text_frame.text = "4th"
        r_ = box.text_frame.paragraphs[0].runs[0]; r_.font.size = Pt(16); r_.font.bold = True; r_.font.color.rgb = rgb("FFFFFF")
        self.text(s, Inches(1.35), Inches(0.72), Inches(6), Inches(0.5), "The Fourth Sheet", 18, bold=True)
        self.text(s, L, Inches(1.7), Inches(12), Inches(0.4), "SAMPLE DATA · invented figures for demonstration", 11, bold=True, colour="8A7A52")
        t = s.shapes.title
        t.left, t.top, t.width, t.height = L, Inches(2.15), CONTENT_W, Inches(1.2)
        t.text_frame.word_wrap = True
        t.text_frame.text = self.title
        para = t.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT
        for r in para.runs:
            r.font.size, r.font.bold, r.font.name = Pt(34), True, FONT
            r.font.color.rgb = rgb(INK)
        self.text(s, L, Inches(3.45), CONTENT_W, Inches(0.5), self.business, 16, colour=MUTED)
        ih = Emu(int(Pt(14).emu * 1.25 * lines_needed(intro, 14, CONTENT_W) + Inches(0.1)))
        self.text(s, L, Inches(4.05), CONTENT_W, ih, intro, 14)
        self.text(s, L, Inches(4.15) + ih, CONTENT_W, Inches(1.2), "How final is this data?\n" + "\n".join(status_lines), 11, colour=MUTED)
        self.notes(s, f"{self.title}. {self.business}. {intro}")

    def _kpi_cards(self, s, items, top, height):
        n = max(1, len(items))
        gap = Inches(0.25)
        w = (CONTENT_W - gap * (n - 1)) / n
        big = 26 if height >= Inches(1.6) else 22
        for i, it in enumerate(items):
            label, value, note = it[:3]
            tone = it[3] if len(it) > 3 else ""
            x = L + i * (w + gap)
            card = s.shapes.add_shape(1, x, top, w, height)
            card.fill.solid(); card.fill.fore_color.rgb = rgb("F7FAF7"); card.line.color.rgb = rgb(LINE)
            self.text(s, x + Inches(0.15), top + Inches(0.08), w - Inches(0.3), Inches(0.42), label, 11, colour=MUTED)
            colour = NEG if (is_neg(value) or tone == "bad") else WARN if tone == "warn" else BRAND
            self.text(s, x + Inches(0.15), top + Inches(0.48), w - Inches(0.3), Inches(0.55), value, big, bold=True, colour=colour)
            self.text(s, x + Inches(0.15), top + height - Inches(0.42), w - Inches(0.3), Inches(0.36), note, 10, colour=MUTED)

    def kpi_slide(self, heading, items, sub=None, notes=None):
        s, top = self.slide(heading, sub, notes or (heading + ": " + "; ".join(f"{a} {b}" for a, b, *_ in items) + "."))
        self._kpi_cards(s, items, top + Inches(0.3), Inches(1.9))
        return s

    def _bar_chart(self, s, x, y, w, h, labels, values, fmt="pct1", below=None, order_by=None, bad=None, warn=None,
                   series_name="Value", target=None, breakeven=None, target_label=None, breakeven_label=None, font=11, keep_order=False):
        by = order_by or values
        order = list(range(len(values)))[::-1] if keep_order else sorted(range(len(values)), key=lambda i: by[i])   # PowerPoint draws bottom-up: largest (or first) on top
        cd = CategoryChartData()
        cd.categories = [labels[i] for i in order]
        cd.add_series(series_name, [values[i] for i in order], number_format=FMT.get(fmt, "General"))
        gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, x, y, w, h, cd)
        ch = gf.chart
        ch.has_title = False                    # writes autoTitleDeleted: no "Value" title
        ch.has_legend = False
        lo = min(0, *values) * 1.15
        hi = max(0, *values, *(t for t in (target, breakeven) if t is not None)) * 1.15 or 1
        va = ch.value_axis
        va.minimum_scale, va.maximum_scale = lo, hi
        va.visible = False
        va.has_major_gridlines = False
        ch.category_axis.tick_labels.font.size = Pt(font)
        ch.category_axis.format.line.fill.background()
        plot = ch.plots[0]
        plot.gap_width = 60
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format, dl.number_format_is_linked = FMT.get(fmt, "General"), False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(font - 1)
        ser = plot.series[0]
        for k, i in enumerate(order):
            pt = ser.points[k]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(NEG if (bad and bad[i]) else WARN if (warn and warn[i]) else GREY if (below and below[i]) else GREEN)
        # lay the plot area out exactly, so a value maps to a known x and the target lines can be drawn on top
        px, py, pw, ph = 0.30, 0.03, 0.58, 0.88
        self._manual_layout(ch, px, py, pw, ph)
        xv = lambda v: x + int(w * (px + pw * (v - lo) / (hi - lo)))
        for v, lab, colour, dash in ((target, target_label, INK, MSO_LINE_DASH_STYLE.DASH), (breakeven, breakeven_label, NEG, MSO_LINE_DASH_STYLE.ROUND_DOT)):
            if v is None:
                continue
            xx = xv(v)
            y0, y1 = y + int(h * py), y + int(h * (py + ph))
            ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, xx, y0, xx, y1)
            ln.line.color.rgb = rgb(colour); ln.line.width = Pt(1.5); ln.line.dash_style = dash
            off = Inches(0.26) if colour == NEG else 0
            self.text(s, xx - Inches(1.2), y1 + off, Inches(2.4), Inches(0.26), lab, 10, bold=True, colour=colour, align=PP_ALIGN.CENTER)
        return ch

    @staticmethod
    def _manual_layout(ch, x, y, w, h):
        plot_area = ch._chartSpace.find(qn("c:chart")).find(qn("c:plotArea"))
        old = plot_area.find(qn("c:layout"))
        if old is not None:
            plot_area.remove(old)
        lay = plot_area.makeelement(qn("c:layout"), {})
        ml = lay.makeelement(qn("c:manualLayout"), {})
        for tag, val in (("layoutTarget", "inner"), ("xMode", "edge"), ("yMode", "edge"), ("x", x), ("y", y), ("w", w), ("h", h)):
            e = ml.makeelement(qn(f"c:{tag}"), {"val": str(val)})
            ml.append(e)
        lay.append(ml)
        plot_area.insert(0, lay)

    def bar_slide(self, heading, labels, values, fmt="pct1", target=None, below=None, sub=None, order_by=None, bad=None, warn=None,
                  series_name=None, target_value=None, breakeven=None, notes=None):
        """target: the label text (e.g. '35.0%'); target_value / breakeven: the values, drawn as lines."""
        s, top = self.slide(heading, sub, notes)
        tv = target_value if target_value is not None else (float(target.rstrip("%")) if isinstance(target, str) and target.rstrip("%").replace(".", "").isdigit() else None)
        self._bar_chart(s, L, top + Inches(0.05), CONTENT_W, H - top - Inches(1.2), labels, values, fmt, below, order_by, bad, warn,
                        series_name or heading.split(" · ")[0], tv, breakeven, f"Target {target}" if target else (f"Target {tv}" if tv is not None else None),
                        f"Break-even {breakeven:.1f}%" if breakeven is not None else None, font=12)
        if below and any(below):
            self.text(s, L, H - Inches(0.72), Inches(6), Inches(0.28), "Grey = below target", 10, colour=MUTED)
        return s

    def _area_chart(self, s, x, y, w, h, labels, values, fmt="money0", series_name="Value", target=None, target_label=None, font=10):
        pts = [(lab, v) for lab, v in zip(labels, values) if v is not None]
        cd = CategoryChartData()
        cd.categories = [lab for lab, _ in pts]
        cd.add_series(series_name, [v for _, v in pts], number_format=FMT.get(fmt, "General"))
        gf = s.shapes.add_chart(XL_CHART_TYPE.AREA, x, y, w, h, cd)
        ch = gf.chart
        ch.has_title = False
        ch.has_legend = False
        vals = [v for _, v in pts]
        lo = min(0, *vals, *( [target] if target is not None else []))
        hi = max(0, *vals, *([target] if target is not None else []))
        span = (hi - lo) or 1
        raw = span / 5
        mag = 10 ** math.floor(math.log10(raw))
        step = next(m * mag for m in (1, 2, 2.5, 5, 10) if m * mag >= raw)      # a round step: 10, 20, 25, 50…
        lo = math.floor(lo / step) * step if lo < 0 else 0
        hi = math.ceil(hi * 1.04 / step) * step
        va = ch.value_axis
        va.minimum_scale, va.maximum_scale, va.major_unit = lo, hi, step
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = rgb("E6EAE7")
        va.tick_labels.number_format, va.tick_labels.number_format_is_linked = FMT.get(fmt, "General"), False
        va.tick_labels.font.size = Pt(font)
        va.format.line.fill.background()
        ca = ch.category_axis
        ca.tick_labels.font.size = Pt(font)
        skip = max(1, math.ceil(len(pts) / 9))
        cax = ca._element                       # show every nth label so dates never crowd (c:tickLblSkip, in schema order)
        for old in cax.findall(qn("c:tickLblSkip")):
            cax.remove(old)
        el = cax.makeelement(qn("c:tickLblSkip"), {"val": str(skip)})
        after = cax.find(qn("c:lblOffset")) if cax.find(qn("c:lblOffset")) is not None else cax.find(qn("c:lblAlgn"))
        if after is not None:
            after.addnext(el)
        else:
            nm = cax.find(qn("c:noMultiLvlLbl"))
            nm.addprevious(el) if nm is not None else cax.append(el)
        ser = ch.plots[0].series[0]
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(GREEN)
        sf = ser.format.fill._xPr.find(qn("a:solidFill"))           # light fill (75% transparent) with a strong line, like the website
        clr = sf.find(qn("a:srgbClr"))
        clr.append(clr.makeelement(qn("a:alpha"), {"val": "25000"}))
        ser.format.line.color.rgb = rgb(GREEN)
        ser.format.line.width = Pt(2.5)
        px, py, pw, ph = 0.11, 0.05, 0.86, 0.80
        self._manual_layout(ch, px, py, pw, ph)
        if target is not None:
            yy = y + int(h * (py + ph * (hi - target) / (hi - lo)))
            ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x + int(w * px), yy, x + int(w * (px + pw)), yy)
            ln.line.color.rgb = rgb(INK); ln.line.width = Pt(1.5); ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            self.text(s, x + int(w * (px + pw)) - Inches(6), yy - Inches(0.3), Inches(6), Inches(0.26), target_label or "Target", 10, bold=True, colour=INK, align=PP_ALIGN.RIGHT)
        return ch

    def area_slide(self, heading, labels, values, fmt="money0", sub=None, series_name=None, target=None, target_label=None, notes=None):
        s, top = self.slide(heading, sub, notes)
        self._area_chart(s, L, top + Inches(0.05), CONTENT_W, H - top - Inches(0.9), labels, values, fmt, series_name or heading.split(" · ")[0], target, target_label)
        return s

    def combo_slide(self, heading, kpis=None, chart=None, blocks=None, sub=None, notes=None, table=None):
        """One idea per slide: headline numbers (label, value, note[, tone]) across the top, the chart in the middle,
        and the notes ([(title, text)], side by side) underneath. Each part only takes the room it needs."""
        s, top = self.slide(heading, sub, notes)
        y = top + Inches(0.05)
        bottom = H - Inches(0.55)
        if kpis:
            self._kpi_cards(s, kpis, y, Inches(1.3))
            y += Inches(1.45)
        nb = len(blocks or [])
        text_h = 0
        if blocks:
            bw = (CONTENT_W - Inches(0.3) * (nb - 1)) / nb
            text_h = max(Inches(0.4) + Emu(int(Pt(12).emu * 1.3 * lines_needed(body, 12, bw - Inches(0.4)))) for _, body in blocks) + Inches(0.15)
            text_h = min(text_h, Inches(2.6))
        if chart:
            ch_h = bottom - y - (text_h + Inches(0.15) if blocks else 0)
            kind = chart.pop("kind")
            if kind == "bars":
                ch_h = min(ch_h, Inches(0.42) * len(chart["labels"]) + Inches(0.8))
                self._bar_chart(s, L, y, CONTENT_W, ch_h, font=11, **chart)
                y += ch_h + (Inches(0.55) if chart.get("breakeven") is not None else Inches(0.35) if chart.get("target") is not None else Inches(0.1))
            else:
                self._area_chart(s, L, y, CONTENT_W, ch_h, **chart)
                y += ch_h + Inches(0.1)
        if blocks:
            bw = (CONTENT_W - Inches(0.3) * (nb - 1)) / nb
            for i, (t, body) in enumerate(blocks):
                x = L + i * (bw + Inches(0.3))
                h = min(text_h, bottom - y)
                card = s.shapes.add_shape(1, x, y, bw, h)
                card.fill.solid(); card.fill.fore_color.rgb = rgb("EEF5F0" if i else "FFFFFF"); card.line.color.rgb = rgb(LINE)
                self.text(s, x + Inches(0.15), y + Inches(0.06), bw - Inches(0.3), Inches(0.3), t.upper(), 10, bold=True, colour=MUTED)
                self.text(s, x + Inches(0.15), y + Inches(0.36), bw - Inches(0.3), h - Inches(0.4), body, 12)
            y += text_h + Inches(0.15)
        if table:
            title, head, rows, tones = (list(table) + [None])[:4]
            self.text(s, L, y, CONTENT_W, Inches(0.3), title.upper(), 10, bold=True, colour=MUTED)
            self._table(s, L, y + Inches(0.3), head, rows, row_h=Inches(0.27), size=10, tones=tones)
        return s

    def content_top(self, heading, sub=None):
        """Where the free area starts under a title (and subtitle), as slide() lays it out."""
        size = 24 if len(heading) <= 70 else 20
        y = Inches(0.42) + Emu(int(Pt(size).emu * 1.2 * min(3, lines_needed(heading, size, CONTENT_W)) + Inches(0.12))) + Inches(0.06)
        if sub:
            y += Emu(int(Pt(12).emu * 1.25 * lines_needed(sub, 12, CONTENT_W) + Inches(0.08))) + Inches(0.06)
        return y

    def fits_table(self, used_top, kpis, blocks, nrows):
        """Would a table of nrows fit under the headline numbers and notes on one combo slide?"""
        y = used_top + Inches(0.05) + (Inches(1.45) if kpis else 0)
        if blocks:
            nb = len(blocks)
            bw = (CONTENT_W - Inches(0.3) * (nb - 1)) / nb
            y += min(Inches(2.6), max(Inches(0.4) + Emu(int(Pt(12).emu * 1.3 * lines_needed(b, 12, bw - Inches(0.4)))) for _, b in blocks) + Inches(0.15)) + Inches(0.15)
        return y + Inches(0.3) + Inches(0.27) * (nrows + 1) < H - Inches(0.5)

    def _table(self, s, x, y, head, rows, row_h=Inches(0.34), size=10.5, tones=None):
        ncol = len(head) if head else max(len(r_) for r_ in rows)
        numeric = [c > 0 and all(looks_numeric(r_[c]) for r_ in rows if c < len(r_)) for c in range(ncol)]
        hrow = 1 if head else 0
        shape = s.shapes.add_table(len(rows) + hrow, ncol, x, y, CONTENT_W, row_h * (len(rows) + hrow))
        t = shape.table
        t.first_row = bool(head)
        if head:
            for c, h_ in enumerate(head):
                cell = t.cell(0, c)
                cell.text = str(h_)
                cell.fill.solid(); cell.fill.fore_color.rgb = rgb(BRAND)
                para = cell.text_frame.paragraphs[0]
                if para.runs:
                    para.runs[0].font.size, para.runs[0].font.bold = Pt(size + 0.5), True
                    para.runs[0].font.color.rgb = rgb("FFFFFF")
                para.alignment = PP_ALIGN.RIGHT if numeric[c] else PP_ALIGN.LEFT
        for r, row_ in enumerate(rows, hrow):
            row_ = (list(row_) + [""] * ncol)[:ncol]
            for c, val in enumerate(row_):
                cell = t.cell(r, c)
                cell.text = "" if val is None else str(val)
                cell.fill.solid(); cell.fill.fore_color.rgb = rgb("FFFFFF" if (r - hrow) % 2 == 0 else "F7FAF7")
                para = cell.text_frame.paragraphs[0]
                tone = tones[r - hrow][c] if tones and r - hrow < len(tones) and c < len(tones[r - hrow]) else ""
                if para.runs:
                    para.runs[0].font.size = Pt(size)
                    para.runs[0].font.color.rgb = rgb(NEG if (is_neg(val) or tone == "bad") else WARN if tone == "warn" else BRAND if tone == "good" else INK)
                    if tone or str(row_[0]).startswith("Total"):
                        para.runs[0].font.bold = True
                para.alignment = PP_ALIGN.RIGHT if numeric[c] else PP_ALIGN.LEFT
        return t

    def table_slides(self, heading, head, rows, sub=None, money_cols=None, notes=None):
        """head=None: no header row (e.g. label / value lists). Text columns are left-aligned, numbers right-aligned."""
        ncol = len(head) if head else max(len(r_) for r_ in rows)
        numeric = [c > 0 and all(looks_numeric(r_[c]) for r_ in rows if c < len(r_)) for c in range(ncol)]
        for start in range(0, max(1, len(rows)), ROWS_PER_SLIDE):
            chunk = rows[start:start + ROWS_PER_SLIDE]
            more = f" ({start // ROWS_PER_SLIDE + 1} of {-(-len(rows) // ROWS_PER_SLIDE)})" if len(rows) > ROWS_PER_SLIDE else ""
            s, top = self.slide(heading + more, sub, notes or (heading + more + (". " + sub if sub else "")))
            hrow = 1 if head else 0
            shape = s.shapes.add_table(len(chunk) + hrow, ncol, L, top + Inches(0.05), CONTENT_W, Inches(0.34) * (len(chunk) + hrow))
            t = shape.table
            t.first_row = bool(head)
            if not head:            # label / value: the label column wider
                t.columns[0].width = Emu(int(CONTENT_W * 0.32))
                for c in range(1, ncol):
                    t.columns[c].width = Emu(int(CONTENT_W * 0.68 / (ncol - 1)))
            if head:
                for c, h_ in enumerate(head):
                    cell = t.cell(0, c)
                    cell.text = str(h_)
                    cell.fill.solid(); cell.fill.fore_color.rgb = rgb(BRAND)
                    para = cell.text_frame.paragraphs[0]
                    if para.runs:
                        para.runs[0].font.size, para.runs[0].font.bold = Pt(11), True
                        para.runs[0].font.color.rgb = rgb("FFFFFF")
                    para.alignment = PP_ALIGN.RIGHT if numeric[c] else PP_ALIGN.LEFT
            for r, row_ in enumerate(chunk, hrow):
                row_ = (list(row_) + [""] * ncol)[:ncol]
                for c, val in enumerate(row_):
                    cell = t.cell(r, c)
                    cell.text = "" if val is None else str(val)
                    cell.fill.solid(); cell.fill.fore_color.rgb = rgb("FFFFFF" if (r - hrow) % 2 == 0 else "F7FAF7")
                    para = cell.text_frame.paragraphs[0]
                    if para.runs:
                        para.runs[0].font.size = Pt(10.5)
                        para.runs[0].font.color.rgb = rgb(NEG if is_neg(val) else INK)
                        if str(row_[0]).startswith("Total") or str(row_[0]) in ("Net profit after tax", "Surplus for the month", "Net assets", "Cash at end of month"):
                            para.runs[0].font.bold = True
                    para.alignment = PP_ALIGN.RIGHT if numeric[c] else PP_ALIGN.LEFT

    def text_slide(self, heading, blocks, sub=None, notes=None):
        """blocks: [(title, text)] side by side, each card as tall as its text."""
        s, top = self.slide(heading, sub, notes or (heading + ": " + " ".join(b for _, b in blocks)))
        n = len(blocks)
        w = (CONTENT_W - Inches(0.3) * (n - 1)) / n
        h = max(Inches(0.55) + Emu(int(Pt(16).emu * 1.3 * lines_needed(body, 16, w - Inches(0.5)))) for _, body in blocks)
        h = min(h, H - top - Inches(0.8))
        for i, (t, body) in enumerate(blocks):
            x = L + i * (w + Inches(0.3))
            card = s.shapes.add_shape(1, x, top + Inches(0.15), w, h)
            card.fill.solid(); card.fill.fore_color.rgb = rgb("EEF5F0" if i else "FFFFFF"); card.line.color.rgb = rgb(LINE)
            self.text(s, x + Inches(0.25), top + Inches(0.25), w - Inches(0.5), Inches(0.35), t.upper(), 11, bold=True, colour=MUTED)
            self.text(s, x + Inches(0.25), top + Inches(0.6), w - Inches(0.5), h - Inches(0.5), body, 16)
        return s

    def image_slide(self, heading, png, sub=None, notes=None):
        s, top = self.slide(heading, sub, notes)
        s.shapes.add_picture(str(png), L, top + Inches(0.05), height=H - top - Inches(0.8))
        return s
